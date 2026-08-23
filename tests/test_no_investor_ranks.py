"""SJAIINV-50: KIS investor_rank 제거 — 소비처가 investor_ranks 를 참조하지 않는다.

이 키는 28주간 항상 비어 있었고(엔드포인트 404), 소비처들은 KRX 전종목
집계 폴백으로 돌아왔다. 제거 후에는 값이 채워져 있어도 무시되어야 한다 —
그래야 죽은 1순위 분기가 되살아나 폴백(더 정확한 전종목 집계)을 가리는
회귀를 막는다.
"""
import openpyxl
import pandas as pd

from modules import database, exporter, reporter

_RANKS = {
    m: pd.DataFrame([
        {"티커": "005930", "종목명": "삼성전자",
         "기관순매수": 100.0, "외국인순매수": 200.0,
         "기관순매수금액": 1111.0, "외국인순매수금액": 2222.0},
    ])
    for m in ["KOSPI", "KOSDAQ"]
}


def test_build_market_data_ignores_investor_ranks():
    processed = {
        "market_info": {"KOSPI_index": {"종가": 6900.0}, "KOSDAQ_index": {"종가": 800.0}},
        "investor_ranks": _RANKS,
    }
    result = database._build_market_data(processed)
    for market in ["KOSPI", "KOSDAQ"]:
        assert "weekly_inst_net" not in result[market]
        assert "weekly_foreign_net" not in result[market]


def test_write_market_sheet_uses_sector_aggregate_not_investor_ranks():
    sector = pd.DataFrame([
        {"섹터": "반도체", "1주외국인매매합(억)": 500.0, "1주기관매매합(억)": 300.0},
        {"섹터": "은행", "1주외국인매매합(억)": -100.0, "1주기관매매합(억)": 50.0},
    ])
    processed = {
        "market_info": {"KOSPI_index": {"종가": 6900.0}, "KOSDAQ_index": {"종가": 800.0}},
        "investor_ranks": _RANKS,
        "kospi_sector": sector,
        "kosdaq_sector": sector,
    }
    ws = openpyxl.Workbook().active
    exporter._write_market_sheet(ws, processed)
    # 3행 = KOSPI: 외국인(G)·기관(H) 주간순매수는 섹터 합산이어야 한다
    assert ws.cell(row=3, column=7).value == 400.0   # 500 - 100
    assert ws.cell(row=3, column=8).value == 350.0   # 300 + 50


def test_report_context_has_no_top30_lines():
    processed = {
        "base_date": "20260821",
        "market_info": {"KOSPI_index": {"종가": 6900.0}, "KOSDAQ_index": {"종가": 800.0}},
        "investor_ranks": _RANKS,
    }
    ctx = reporter._build_data_context(processed, {}, {})
    assert "Top30" not in ctx["market_summary"]


def test_processor_ignores_investor_ranks_and_drops_key():
    """네 번째 소비처(processor). KIS 병합이 살아 있으면 naver 원본에 이미 있는
    1주기관매매와 rename 이 충돌해 중복 컬럼을 만들 수 있다 — 무시가 정답."""
    raw = {
        "kospi": pd.DataFrame({
            "티커": ["005930"], "종목명": ["삼성전자"],
            "시가총액(억)": [4000000.0], "시장": ["KOSPI"],
            "1주기관매매": [999.0], "1주외국인매매": [888.0],
        }),
        "kosdaq": pd.DataFrame(),
        "market_info": {}, "base_date": "20260821",
        "krx_flows": pd.DataFrame(), "flow_source": "naver",
        "investor_ranks": _RANKS,
    }
    from modules import processor
    out = processor.process(raw)
    df = out["kospi"]
    assert list(df.columns).count("1주기관매매") == 1
    assert df.set_index("티커").loc["005930", "1주기관매매"] == 999.0  # KIS 값 미반영
    assert "investor_ranks" not in out
